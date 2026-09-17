"""Every U.S. House and Senate candidate's primary and general votes, out of
the FEC's biennial "Federal Elections" workbooks, as JSON for load.mjs.

    python3 scripts/elections/fec_primaries.py ~/Code/govblock-data/elections/fec

Reads federalelections<year>.xlsx (2016–2022) or .xls (2008–2014, needs
xlrd) in that folder and writes fec-candidates-<year>.json beside each. The
workbooks change their column names and sheet layout from edition to edition
("PRIMARY" in 2008, "PRIMARY VOTES" later; House and Senate on one sheet
until 2012), so the candidate table is found by its header, not its name.

A primary cell holds votes, or a mark the FEC explains in the row's footnote:
"Unopposed" (no primary was counted), "*" (nominated by a party convention),
"#" (another footnote). The marks are kept as they are; load.mjs decides what
a race without counted votes can say.
"""

import json
import os
import sys

YEARS = (2008, 2010, 2012, 2014, 2016, 2018, 2020, 2022)

SKIP_NAMES = {"scattered", "scattering", "district votes:", "total state votes:", "write-in", "write-ins", "none of these candidates", "blank votes", "void votes", "over votes", "under votes", "total votes:"}

ALIASES = {
    "state": ("STATE ABBREVIATION",),
    "district": ("DISTRICT", "D"),
    "fec_id": ("FEC ID", "FEC ID#"),
    "incumbent": ("(I)", "(I) INCUMBENT INDICATOR", "(I) INCUMBENT INDICATOR", "INCUMBENT INDICATOR (I)"),
    "name": ("CANDIDATE NAME", "CANDIDATE NAME (LAST, FIRST)"),
    "party": ("PARTY",),
    "primary": ("PRIMARY VOTES", "PRIMARY"),
    "general": ("GENERAL VOTES", "GENERAL"),
    "winner": ("GE WINNER INDICATOR",),
}


def sheets(path):
    """Yields (sheet name, row iterator) for either workbook format."""
    if path.endswith(".xls"):
        import xlrd

        wb = xlrd.open_workbook(path)
        for sh in wb.sheets():
            yield sh.name, (sh.row_values(r) for r in range(sh.nrows))
    else:
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        for name in wb.sheetnames:
            yield name, wb[name].iter_rows(values_only=True)


def text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return int(value)
    return None


def columns(header):
    upper = [text(c).upper() for c in header]
    found = {}
    for key, names in ALIASES.items():
        for i, h in enumerate(upper):
            if h in names:
                found[key] = i
                break
    return found


def extract(path, year):
    out = []
    for sheet, rows in sheets(path):
        # The House and Senate results tables; not the special elections held
        # off-cycle, the presidential tables or the party summaries.
        lower = sheet.lower()
        if not ("house" in lower or "senate" in lower) or "by party" in lower or "special" in lower or "gains" in lower:
            continue
        cols = None
        for row in rows:
            row = list(row)
            if cols is None:
                found = columns(row)
                if {"state", "name", "primary", "general", "district"} <= found.keys():
                    cols = found
                continue
            get = lambda key: row[cols[key]] if key in cols and cols[key] < len(row) else None
            state = text(get("state"))
            name = text(get("name"))
            party = text(get("party"))
            if len(state) != 2 or not name or name.lower() in SKIP_NAMES or party.lower().startswith("combined"):
                continue
            raw_district = text(get("district"))
            if not raw_district or raw_district.lower() == "n/a":
                continue
            senate = raw_district.upper().startswith("S")
            if senate:
                office, district, label = "US SENATE", "", raw_district
            else:
                head = raw_district.split(" ")[0]
                office = "US HOUSE"
                district = str(int(head)) if head.isdigit() else ("0" if head.upper() in ("AL", "AT-LARGE") else head.upper())
                label = raw_district
            primary_cell = get("primary")
            fec_id = text(get("fec_id"))
            out.append(
                {
                    "year": year,
                    "office": office,
                    "state": state,
                    "district": district,
                    "label": label,
                    "fec_id": fec_id if fec_id and fec_id.lower() != "n/a" else None,
                    "name": name,
                    "party": party,
                    "incumbent": bool(text(get("incumbent"))),
                    "primary": number(primary_cell),
                    "primary_mark": text(primary_cell) if isinstance(primary_cell, str) and text(primary_cell) else None,
                    "general": number(get("general")),
                    "winner": text(get("winner")) == "W",
                }
            )
    return out


def main():
    folder = os.path.expanduser(sys.argv[1] if len(sys.argv) > 1 else "~/Code/govblock-data/elections/fec")
    for year in YEARS:
        path = next((p for p in (os.path.join(folder, f"federalelections{year}.xlsx"), os.path.join(folder, f"federalelections{year}.xls")) if os.path.exists(p)), None)
        if not path:
            print(f"{year}: no workbook")
            continue
        rows = extract(path, year)
        with open(os.path.join(folder, f"fec-candidates-{year}.json"), "w") as f:
            json.dump(rows, f)
        house = sum(1 for r in rows if r["office"] == "US HOUSE")
        print(f"{year}: {len(rows)} candidate rows ({house} House), {sum(1 for r in rows if r['primary'] is not None)} with primary votes, {sum(1 for r in rows if r['winner'])} marked winners")


if __name__ == "__main__":
    main()
